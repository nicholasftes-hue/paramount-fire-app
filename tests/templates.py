"""Verify editable documents and PDF field save/reopen using fictional values."""
import io
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
from docx import Document
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
results = []
for kind in ['Quotation', 'Invoice']:
    word_path = ROOT / 'templates' / f'Paramount_{kind}_Editable.docx'
    with ZipFile(word_path) as archive:
        assert archive.testzip() is None
        assert 'word/document.xml' in archive.namelist()
        assert not any('vbaProject' in name for name in archive.namelist())
        assert b'documentProtection' not in archive.read('word/settings.xml')
    doc = Document(word_path)
    assert len(doc.tables) == 2
    assert len(doc.tables[1].rows) == 6
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'DRAFT TEMPLATE' in text and 'no automatic calculation' in text
    assert kind.upper() in text and 'Not a fire safety certificate' in text
    doc.tables[0].cell(0, 0).text = 'TEST ONLY - fictional company'
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    assert Document(stream).tables[0].cell(0, 0).text == 'TEST ONLY - fictional company'
    pdf_path = ROOT / 'templates' / f'Paramount_{kind}_Fillable.pdf'
    reader = PdfReader(pdf_path, strict=True)
    assert len(reader.pages) == 1
    assert not reader.is_encrypted
    assert '/JavaScript' not in str(reader.trailer) and '/OpenAction' not in reader.trailer['/Root']
    fields = reader.get_fields()
    assert len(fields) == (43 if kind == 'Quotation' else 45), len(fields)
    required = ['company_address','customer_address','site_address','document_number','issue_date',
                'until_date','vat_registration','subtotal','tax','total','terms']
    assert all(name in fields for name in required)
    for name, field in fields.items():
        assert field['/FT'] == '/Tx' and not (int(field.get('/Ff', 0)) & 1), name
        assert not field.get('/V'), name
        assert field.get('/TU'), name
    text = reader.pages[0].extract_text()
    assert 'DRAFT TEMPLATE' in text and 'Not a fire safety certificate' in text
    widgets = reader.pages[0]['/Annots']
    assert len(widgets) == len(fields)
    for widget in widgets:
        x1, y1, x2, y2 = widget.get_object()['/Rect']
        assert 0 <= x1 < x2 <= 596 and 0 <= y1 < y2 <= 842
    writer = PdfWriter(clone_from=reader)
    values = {name: f'TEST {i}' for i, name in enumerate(fields)}
    writer.update_page_form_field_values(writer.pages[0], values, auto_regenerate=False)
    with TemporaryDirectory() as directory:
        output = Path(directory) / 'test-filled.pdf'
        writer.write(output)
        reopened = PdfReader(output, strict=True)
        for name, value in values.items():
            assert reopened.get_fields()[name]['/V'] == value, name
        for widget in reopened.pages[0]['/Annots']:
            assert '/AP' in widget.get_object(), 'Saved field needs a visible appearance'
    results.append({'document':kind,'word_edit_save_reopen':'passed','pdf_fields':len(fields),
                    'pdf_fill_save_reopen_all_fields':'passed'})
print(json.dumps(results, indent=2))
