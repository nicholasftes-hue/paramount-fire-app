"""Independent strict PDF parsing and editable Word table round trips."""
from pathlib import Path
from docx import Document
from pypdf import PdfReader
for kind,markers in [('extinguisher',['EX-1','EX-2','Fictional damaged hose','SER-1']),('structured-fire',['DEV-1','L1:001','Smoke detector','SYSTEM / PANEL / BATTERIES'])]:
    for width in [320,360,412,800]:
        stem=Path('test-results')/f'{kind}-{width}'
        pdf=PdfReader(stem.with_suffix('.pdf'),strict=True)
        text='\n'.join(p.extract_text() for p in pdf.pages)
        for marker in markers: assert marker in text, (stem,marker)
        assert all('NOT A COMPLIANCE CERTIFICATE' in p.extract_text() for p in pdf.pages)
        doc=Document(stem.with_suffix('.docx'))
        text='\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
        for marker in markers: assert marker in text,(stem,marker)
        doc.tables[0].cell(0,1).text+=' FICTIONAL OFFICE EDIT'
        target=stem.with_name(stem.name+'-edited').with_suffix('.docx');doc.save(target)
        assert 'FICTIONAL OFFICE EDIT' in Document(target).tables[0].cell(0,1).text
        print('PASS strict PDF and Word table edit/save/reopen:',stem)
