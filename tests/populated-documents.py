"""Verify actual saved-record Word exports independently of the JS ZIP writer."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree
from docx import Document
from pypdf import PdfReader

for kind in ('service', 'quote'):
    path = Path('test-results') / f'populated-{kind}.docx'
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        assert len(archive.namelist()) == 5
        for name in archive.namelist():
            ElementTree.fromstring(archive.read(name))
    doc = Document(path)
    text = '\n'.join([p.text for p in doc.paragraphs]+[c.text for t in doc.tables for row in t.rows for c in row.cells])
    assert 'Reference' in text and 'Revision' in text
    assert 'NOT A COMPLIANCE CERTIFICATE' in doc.sections[0].footer.paragraphs[0].text
    if kind == 'service':
        assert '24.125' in text and 'battery' in text.lower() and 'Zoë & <test> Ω' in text
    else:
        assert 'Total: £0.36' in ' '.join(text.split()) and 'Service & parts' in text
    doc.paragraphs[-1].add_run(' FICTIONAL EDIT CHECK')
    output = path.with_name(f'reopened-{kind}.docx')
    doc.save(output)
    assert 'FICTIONAL EDIT CHECK' in Document(output).paragraphs[-1].text
    print(f'PASS {kind}: Word ZIP, XML, saved values, footer, edit/save/reopen')
pdf = PdfReader('test-results/populated-quote.pdf', strict=True)
text = '\n'.join(p.extract_text() for p in pdf.pages)
assert 'Total: £0.36' in ' '.join(text.split()) and 'Service & parts' in text
assert all('NOT A COMPLIANCE CERTIFICATE' in p.extract_text() for p in pdf.pages)
print('PASS populated quote PDF: independent strict parsing and totals')
