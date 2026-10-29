"""Independent parsing and edit/save/reopen checks for the added certificate types."""
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from pypdf import PdfReader
for kind in ['extinguisher','installation']:
    stem=Path('test-results')/('premium-'+kind)
    reader=PdfReader(stem.with_suffix('.pdf'),strict=True)
    text='\n'.join(page.extract_text() for page in reader.pages)
    for expected in ['DRAFT SIGNATURE ACKNOWLEDGEMENTS','Fictional customer','Fictional repair','Revision: 1']:
        assert expected in text,(kind,expected)
    path=stem.with_suffix('.docx')
    with ZipFile(path) as archive:
        footer=archive.read('word/footer1.xml').decode()
        assert 'NUMPAGES' in footer and 'PAGE' in footer and 'Revision: 1' in footer
    document=Document(path)
    target=next(cell for table in document.tables for row in table.rows for cell in row.cells if 'Fictional repair' in cell.text)
    target.text='Office edited remedial advice'
    edited=stem.with_name(stem.name+'-office-edited').with_suffix('.docx')
    document.save(edited)
    assert any('Office edited remedial advice' in cell.text for table in Document(edited).tables for row in table.rows for cell in row.cells)
    print('PASS',kind,'strict PDF, revision/page footer and Word table edit/save/reopen')
