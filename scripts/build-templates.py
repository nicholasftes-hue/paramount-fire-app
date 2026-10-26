"""Generate blank, editable commercial templates; no app/customer data input.

Requires python-docx and reportlab. These are recreated templates, not recovered
originals. All commercial details and totals require review before issue.
"""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import A4

ROOT = Path(__file__).resolve().parents[1] / 'templates'
ROOT.mkdir(exist_ok=True)
NAVY = '#18334b'
RED = '#b3262c'
NOTICE = 'DRAFT TEMPLATE - Complete and review before issue. Not a fire safety certificate.'
TERMS = 'Enter agreed scope, exclusions, payment terms and applicable tax treatment. Totals are entered manually; no automatic calculation.'


def word(kind):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.55)
    section.left_margin = section.right_margin = Inches(.6)
    style = doc.styles['Normal']
    style.font.name, style.font.size = 'Calibri', Pt(10)
    style.paragraph_format.space_after = Pt(5)
    title = doc.add_paragraph('PARAMOUNT FIRE PROTECTION')
    title.runs[0].font.size = Pt(21)
    title.runs[0].font.bold = True
    title.runs[0].font.color.rgb = RGBColor.from_string('18334B')
    doc.add_heading(kind.upper(), 1)
    doc.add_paragraph(NOTICE)
    detail = doc.add_table(rows=0, cols=2)
    detail.style = 'Table Grid'
    date_label = 'Valid until' if kind == 'Quotation' else 'Payment due'
    for a, b in [
        ('Company address: [Enter company address]', f'{kind} number: [Enter reference]'),
        ('Telephone / email: [Enter contact details]', 'Issue date: [Enter date]'),
        ('Company registration: [Enter if applicable]', f'{date_label}: [Enter date]'),
        ('VAT registration: [Enter if applicable]', 'Job / order reference: [Enter reference]'),
        ('Customer / billing address: [Enter details]', 'Site / work address: [Enter details]'),
        ('Customer contact: [Enter authorised contact]', 'Prepared by: [Enter name]'),
    ]:
        cells = detail.add_row().cells
        cells[0].text, cells[1].text = a, b
    doc.add_heading('Scope and itemised charges (GBP)', 2)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    for cell, label in zip(table.rows[0].cells, ['Description / scope', 'Qty', 'Unit GBP', 'VAT %', 'Net GBP']):
        cell.text = label
    for _ in range(5):
        for cell, value in zip(table.add_row().cells, ['[Describe work / materials]', '[Qty]', '[Rate]', '[VAT]', '[Net]']):
            cell.text = value
    doc.add_paragraph('Net subtotal GBP: [Enter]     VAT total GBP: [Enter]     Total GBP: [Enter]')
    if kind == 'Invoice':
        doc.add_paragraph('Paid GBP: [Enter]     Balance due GBP: [Enter]')
    doc.add_heading('Terms, exclusions and notes', 2)
    doc.add_paragraph(TERMS)
    doc.add_paragraph('[Enter agreed terms, limitations and follow-up arrangements]')
    if kind == 'Quotation':
        doc.add_paragraph('Acceptance: [Name / signature]     Date: [Enter]')
        doc.add_paragraph('Acceptance agrees commercial scope only; it does not certify completed work.')
    else:
        doc.add_paragraph('Payment instructions: [Enter verified payment details and payment reference]')
    section.footer.paragraphs[0].text = 'Review company details, prices, VAT, dates and authorisation before sending. No compliance or competency sign-off is provided.'
    doc.core_properties.title = f'Paramount {kind} - editable draft template'
    doc.core_properties.author = 'Paramount Fire Protection'
    doc.save(ROOT / f'Paramount_{kind}_Editable.docx')


def pdf(kind):
    c = canvas.Canvas(str(ROOT / f'Paramount_{kind}_Fillable.pdf'), pagesize=A4)
    c.setTitle(f'Paramount {kind} - fillable draft template')
    c.setAuthor('Paramount Fire Protection')
    width, height = A4
    c.setFillColor(HexColor(NAVY))
    c.rect(0, height - 93, width, 93, fill=1, stroke=0)
    c.setFillColor(white)
    c.setFont('Helvetica-Bold', 20)
    c.drawString(32, height - 35, 'PARAMOUNT FIRE PROTECTION')
    c.setFont('Helvetica', 17)
    c.drawString(32, height - 64, kind.upper())
    c.setFillColor(HexColor(RED))
    c.setFont('Helvetica-Bold', 8)
    c.drawString(32, height - 110, NOTICE)

    def text(x, y, value, size=8):
        c.setFillColor(HexColor(NAVY))
        c.setFont('Helvetica', size)
        c.drawString(x, y, value)

    def field(name, label, x, y, w, h=20, multiline=False):
        text(x, y + h + 4, label)
        c.acroForm.textfield(
            name=name, tooltip=label, x=x, y=y, width=w, height=h,
            borderWidth=.5, borderColor=HexColor('#91a4b3'),
            fillColor=HexColor('#f4f7fa'), textColor=HexColor(NAVY),
            fontName='Helvetica', fontSize=9, value='',
            fieldFlags='multiline' if multiline else '',
            maxlen=1200 if multiline else 160,
        )

    for name, label, x, y, w in [
        ('company_address','Company address',32,675,258),
        ('document_number',f'{kind} number',305,675,258),
        ('company_contact','Company telephone / email',32,632,258),
        ('issue_date','Issue date',305,632,123),
        ('until_date','Valid until' if kind=='Quotation' else 'Payment due',440,632,123),
        ('company_registration','Company registration (if applicable)',32,589,258),
        ('vat_registration','VAT registration (if applicable)',305,589,258),
        ('customer_address','Customer / billing address',32,531,258),
        ('site_address','Site / work address',305,531,258),
        ('customer_contact','Customer contact',32,488,258),
        ('job_reference','Job / order reference',305,488,123),
        ('prepared_by','Prepared by',440,488,123),
    ]:
        field(name,label,x,y,w,35 if name.endswith('address') else 20,name.endswith('address'))
    text(32,466,'ITEMISED CHARGES - GBP (manual totals)',10)
    columns = [(32,270,'Description / scope','description'),(310,45,'Qty','quantity'),(363,62,'Unit GBP','rate'),(433,52,'VAT %','vat'),(493,70,'Net GBP','net')]
    for i in range(1,6):
        for x,w,label,key in columns:
            # Tooltips retain row and column names for assistive PDF viewers.
            field(f'item_{i}_{key}',f'{label} {i}',x,420-(i-1)*35,w)
    for name,label,x in [('subtotal','Net subtotal GBP',32),('tax','VAT total GBP',213),('total','Total GBP',394)]:
        field(name,label,x,236,169)
    if kind == 'Invoice':
        field('paid','Paid GBP',32,194,169)
        field('balance','Balance due GBP',213,194,169)
        field('payment_reference','Payment reference',394,194,169)
    text(32,177,'Terms / exclusions / notes - complete agreed scope and tax treatment')
    field('terms','Totals are manual. No automatic calculation.',32,112,531,47,True)
    if kind == 'Quotation':
        field('acceptance','Acceptance name (not a digital signature)',32,64,350)
        field('acceptance_date','Acceptance date',394,64,169)
    else:
        field('payment_instructions','Verified payment instructions',32,64,531)
    text(32,42,'Review details, prices, VAT and authorisation before issue. Save using a PDF editor.',7)
    text(32,30,'Commercial template only. No fire safety, system compliance or competency certification.',7)
    c.showPage()
    c.save()


if __name__ == '__main__':
    for kind in ['Quotation','Invoice']:
        word(kind)
        pdf(kind)
        print(f'Created {kind}: editable DOCX and fillable PDF')
